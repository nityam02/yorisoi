"""Build the Yorisoi pitch deck on the AI Builder Cup's prescribed template.

    uv run --no-project --with python-pptx --with pillow python tools/deck/build_deck.py \
        [--prototype-url https://yorisoi-xxxx.a.run.app] [--pdf]

Edits the official template in place (keeps its slides, titles, banner and footer), removes only the
template's instruction slide ("Please download this template..."), and fills every other slide.
Output: tools/deck/out/Yorisoi-Pitch-Deck.pptx (+ .pdf with --pdf, via LibreOffice).
"""

import argparse
import copy
import shutil
import subprocess
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
ASSETS = HERE / "assets"
GITHUB = "https://github.com/nityam02/yorisoi"
VIDEO = "https://github.com/nityam02/yorisoi/releases/download/v1.0/yorisoi-demo.mp4"

# Template fonts (Google Slides names) and a quiet Google palette.
F_REG, F_MED, F_SEMI = "Google Sans Flex", "Google Sans Flex Medium", "Google Sans Flex SemiBold"
INK, INK2, MUTED = RGBColor(0x20, 0x21, 0x24), RGBColor(0x3C, 0x40, 0x43), RGBColor(0x5F, 0x63, 0x68)
BLUE, BLUE_T = RGBColor(0x1A, 0x73, 0xE8), RGBColor(0xE8, 0xF0, 0xFE)
GREY_T, GREY_L = RGBColor(0xF1, 0xF3, 0xF4), RGBColor(0xDA, 0xDC, 0xE0)
RED, RED_T = RGBColor(0xC5, 0x22, 0x1F), RGBColor(0xFC, 0xE8, 0xE6)
GREEN, GREEN_T = RGBColor(0x18, 0x80, 0x38), RGBColor(0xE6, 0xF4, 0xEA)

X0, XW = 0.39, 9.22  # content left edge and width (aligned with the template's title text)
Y0 = 1.42  # content top, below the template's 16 pt slide title


# ---------------------------------------------------------------- text helpers

def _font(run, size, font=F_REG, color=INK, bold=False, link=None):
    run.font.size = Pt(size)
    run.font.name = font
    run.font.bold = bold
    run.font.color.rgb = color
    rpr = run._r.get_or_add_rPr()
    # Google Sans Flex has no Japanese glyphs, so Japanese text uses Noto Sans JP.
    jp = "Noto Sans JP Medium" if font in (F_MED, F_SEMI) else "Noto Sans JP"
    for tag, face in (("a:ea", jp), ("a:cs", font)):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", face)
    if link:
        run.hyperlink.address = link


def _bullet(paragraph, char="•", indent_in=0.17):
    ppr = paragraph._p.get_or_add_pPr()
    ppr.set("marL", str(int(Inches(indent_in))))
    ppr.set("indent", str(-int(Inches(indent_in))))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in ppr.findall(qn(tag)):
            ppr.remove(el)
    bu = etree.SubElement(ppr, qn("a:buChar"))
    bu.set("char", char)


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, name=None, inset=0.0):
    """paras: list of dicts {runs: [(text, opts)], bullet, after, before, align}."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, Inches(inset))
    tf.vertical_anchor = anchor
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = p.get("align", PP_ALIGN.LEFT)
        para.space_after = Pt(p.get("after", 4))
        para.space_before = Pt(p.get("before", 0))
        para.line_spacing = p.get("line", 1.05)
        if p.get("bullet"):
            _bullet(para)
        for t, o in p["runs"]:
            r = para.add_run()
            r.text = t
            _font(r, **o)
    return box


def P(*runs, bullet=False, after=4, before=0, align=PP_ALIGN.LEFT, line=1.05):
    return {"runs": list(runs), "bullet": bullet, "after": after, "before": before, "align": align, "line": line}


def R(t, size=11, font=F_REG, color=INK, bold=False, link=None):
    return (t, {"size": size, "font": font, "color": color, "bold": bold, "link": link})


def box(slide, x, y, w, h, fill=GREY_T, line=None, radius=0.06, shape=MSO_SHAPE.ROUNDED_RECTANGLE, lw=1.0):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.text_frame.text = ""
    return s


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=1.5):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return c


def picture(slide, name, x, y, w):
    pic = slide.shapes.add_picture(str(ASSETS / name), Inches(x), Inches(y), width=Inches(w))
    pic.line.color.rgb = GREY_L
    pic.line.width = Pt(0.75)
    return pic


def card(slide, x, y, w, h, title, body, fill=GREY_T, title_color=INK, size=10.5, title_size=12, bullets=False):
    box(slide, x, y, w, h, fill=fill)
    paras = [P(R(title, title_size, F_SEMI, title_color), after=5)]
    for b in body:
        paras.append(P(R(b, size, F_REG, INK2), bullet=bullets, after=3))
    text(slide, x + 0.14, y + 0.1, w - 0.28, h - 0.18, paras)


# ---------------------------------------------------------------- template helpers

def set_title(slide, new_text, keep_paragraphs=1):
    """Replace the template title text, keep its run formatting, drop prompt paragraphs below it."""
    tb = next(s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip())
    paras = tb.text_frame.paragraphs
    first = paras[0]
    runs = first.runs
    runs[0].text = new_text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    for br in first._p.findall(qn("a:br")):
        first._p.remove(br)
    for p in paras[keep_paragraphs:]:
        p._p.getparent().remove(p._p)
    tb.height = Inches(0.61)
    return tb


def delete_slide(prs, index):
    sld_ids = prs.slides._sldIdLst
    sld = list(sld_ids)[index]
    prs.part.drop_rel(sld.get(qn("r:id")))
    sld_ids.remove(sld)


# ---------------------------------------------------------------- slides

def s_team(slide):
    tb = next(s for s in slide.shapes if s.has_text_frame and "Team Details" in s.text_frame.text)
    values = {"Team name:": "The AI way", "Team leader name:": "Aayushi Sinha",
              "Problem Statement:": "Sustainability & Social Impact"}
    last = None
    for p in tb.text_frame.paragraphs:
        label = p.text.strip()
        if label in values:
            r = p.runs[0]
            r.text = label + " "
            new = copy.deepcopy(r._r)
            new.find(qn("a:t")).text = values[label]
            r._r.addnext(new)
            last = p
    # One more item in the same style: the solution's name.
    extra = copy.deepcopy(last._p)
    last._p.addnext(extra)
    ts = extra.findall(".//" + qn("a:t"))
    ts[0].text = "Solution: "
    ts[1].text = "Yorisoi (寄り添い), an AI copilot for individual evacuation plans"
    ea = ts[1].getparent().find(qn("a:rPr")).find(qn("a:ea"))
    ea.set("typeface", "Noto Sans JP")
    # Drop the trailing empty paragraph so the list keeps its spacing.
    paras = tb.text_frame.paragraphs
    if not paras[-1].text.strip():
        paras[-1]._p.getparent().remove(paras[-1]._p)


def s_brief(slide):
    set_title(slide, "Brief about the idea")
    text(slide, X0, Y0, 5.75, 0.95, [P(R(
        "Yorisoi (寄り添い, “staying by someone’s side”) helps care workers turn one conversation with an older or "
        "disabled resident into a hazard-checked, human-approved individual evacuation plan.", 13, F_MED, INK),
        line=1.1)])
    text(slide, X0, 2.45, 5.75, 2.95, [
        P(R("The problem", 12, F_SEMI, INK), after=4),
        P(R("Japan’s 1,741 municipalities keep registers of residents who need help evacuating. Since 2021 they must "
            "endeavour to make each of them an individual evacuation plan.", 11, F_REG, INK2), bullet=True),
        P(R("Each plan is written by hand: an interview, a hazard check of the home, a suitable shelter, a route "
            "and named helpers.", 11, F_REG, INK2), bullet=True),
        P(R("In the 2011 disaster, people with disability certificates died at 1.43% against 0.78% overall. In Mabi "
            "(2018), 51 people died, most of them elderly, where the hazard map had marked the risk.", 11, F_REG, INK2),
          bullet=True),
        P(R("Yorisoi does the slow parts with Gemini and public data. A person approves every plan.",
            11, F_MED, INK), before=2),
    ])
    box(slide, 6.45, Y0, 3.16, 3.95, fill=BLUE_T)
    text(slide, 6.62, Y0 + 0.12, 2.85, 3.75, [
        P(R("15.1%", 44, F_SEMI, BLUE), after=0, line=0.95),
        P(R("of 6,696,718 people on the registers had a plan on 1 April 2026 (1,014,017).", 11, F_REG, INK), after=10),
        P(R("911 of 1,741", 18, F_SEMI, INK), after=0),
        P(R("municipalities cover 20% of their list or less.", 11, F_REG, INK2), after=10),
        P(R("≈ 29 years", 18, F_SEMI, INK), after=0),
        P(R("to finish at last year’s pace of 198,106 new plans.", 11, F_REG, INK2), after=10),
        P(R("Source: Cabinet Office and Fire and Disaster Management Agency survey, published 29 June 2026. "
            "2011 figures: NHK via Japan Disability Forum.", 8.5, F_REG, MUTED)),
    ])


def s_opportunities(slide):
    set_title(slide, "Opportunities")
    cols = [
        ("How different is it from any of the other existing ideas?", [
            "Municipal tools from NEC and Pasona move the paper form onto a tablet. Staff still interview, read "
            "hazard maps, pick shelters and write.",
            "Consumer apps such as Bracewise and BePrepared make household checklists from US hazard data.",
            "We found no tool that automates the interview, the hazard check and drafting to the official form.",
        ]),
        ("How will it be able to solve the problem?", [
            "It removes the slow steps of each plan instead of digitising the form.",
            "Gemini turns the conversation into facts backed by exact quotes. Public GSI data supplies flood depth "
            "and certified shelters.",
            "Gemini drafts the Cabinet Office form with evidence on every line; rules flag gaps before a supervisor "
            "approves.",
        ]),
        ("USP of the proposed solution", [
            "No quote, no field: every line traces to the conversation, the register or a public-data tool.",
            "Facts come from deterministic tools, never the model. Missing data shows “no data”, never “safe”.",
            "Must-fix checks block approval. Easy-read Japanese, a helper card and translations come out of the "
            "same plan.",
        ]),
    ]
    w, gap = 2.95, 0.185
    for i, (title, body) in enumerate(cols):
        x = X0 + i * (w + gap)
        card(slide, x, Y0, w, 3.95, title, body, fill=BLUE_T if i == 2 else GREY_T,
             title_color=BLUE if i == 2 else INK, size=11.5, bullets=True)


def s_features(slide):
    set_title(slide, "List of features offered by the solution")
    feats = [
        ("Risk-ranked worklist", "Residents ordered by flood depth at home, living alone, mobility and medical devices. "
                                 "Missing information raises priority."),
        ("Hazard brief before the visit", "Flood depth read from the national Hazard Map Portal tile at the home, the "
                                          "nearest GSI shelters certified for floods, and the path checked for deep water."),
        ("Conversation capture", "Record the visit and Gemini transcribes the Japanese audio, or upload a recording, "
                                 "or type the transcript."),
        ("Evidence-locked draft", "Gemini fills the Cabinet Office sample form. Every line cites its source, and quotes "
                                  "are checked against the transcript in code."),
        ("Safety checks", "Rules check shelter certification, two helpers with phone numbers, power for medical devices, "
                          "consent and transport. A Gemini review can only add notes."),
        ("Copies and approval", "Easy-read Japanese, a helper card and a translation. The resident confirms, a "
                                "supervisor approves with a seal; CSV export and an audit log."),
    ]
    w, h, gx, gy = 2.95, 1.86, 0.185, 0.2
    for i, (title, body) in enumerate(feats):
        x = X0 + (i % 3) * (w + gx)
        y = Y0 + (i // 3) * (h + gy)
        card(slide, x, y, w, h, title, [body], fill=GREY_T, size=11.5)


def _step(slide, x, y, w, h, n, title, body, kind):
    fill, line = {"gemini": (BLUE_T, BLUE), "code": (GREY_T, MUTED), "people": (RED_T, RED)}[kind]
    box(slide, x, y, w, h, fill=fill, line=line, lw=1.0)
    dot = box(slide, x + 0.1, y + 0.1, 0.3, 0.3, fill=line, shape=MSO_SHAPE.OVAL)
    text(slide, x + 0.1, y + 0.1, 0.3, 0.3, [P(R(str(n), 10, F_SEMI, RGBColor(0xFF, 0xFF, 0xFF)),
                                              align=PP_ALIGN.CENTER, after=0)], anchor=MSO_ANCHOR.MIDDLE)
    text(slide, x + 0.47, y + 0.08, w - 0.55, 0.36, [P(R(title, 11, F_SEMI, INK), after=0)],
         anchor=MSO_ANCHOR.MIDDLE)
    text(slide, x + 0.12, y + 0.5, w - 0.22, h - 0.56, [P(R(body, 9.5, F_REG, INK2), after=0)])
    return dot


def s_flow(slide):
    set_title(slide, "Process flow diagram or Use-case diagram")
    w, h, gap = 2.06, 1.42, 0.33
    xs = [X0 + i * (w + gap) for i in range(4)]
    y1, y2 = Y0 + 0.05, Y0 + 2.0
    row1 = [("Pick a resident", "From the register, ranked by risk; unknowns rank higher.", "code"),
            ("Hazard brief", "GSI flood depth at home, flood-certified shelters, path check.", "code"),
            ("Conversation", "Recorded and transcribed by Gemini, uploaded, or typed.", "people"),
            ("Extract facts", "Gemini returns facts with exact quotes; unquoted items are dropped.", "gemini")]
    row2 = [("Draft the plan", "Gemini fills the official fields and cites evidence; shelter only from the tool list.",
             "gemini"),
            ("Safety checks", "Rules in code, then a Gemini review that can only add notes.", "code"),
            ("Review and approve", "Care worker edits, resident confirms, supervisor approves. Must-fix checks block.",
             "people"),
            ("Outputs", "Official plan (print, CSV), easy-read copy, helper card, translation.", "code")]
    for i, (t, b, k) in enumerate(row1):
        _step(slide, xs[i], y1, w, h, i + 1, t, b, k)
        if i < 3:
            arrow(slide, xs[i] + w + 0.04, y1 + h / 2, xs[i + 1] - 0.04, y1 + h / 2)
    # Row 2 runs right to left so the flow snakes down from step 4.
    for j, (t, b, k) in enumerate(row2):
        i = 3 - j
        _step(slide, xs[i], y2, w, h, j + 5, t, b, k)
        if j < 3:
            arrow(slide, xs[i] - 0.04, y2 + h / 2, xs[i - 1] + w + 0.04, y2 + h / 2)
    arrow(slide, xs[3] + w / 2, y1 + h + 0.04, xs[3] + w / 2, y2 - 0.04)
    ly = y2 + h + 0.17
    for k, (label, kind) in enumerate([("Gemini", "gemini"), ("Code and public data", "code"), ("People", "people")]):
        fill, line = {"gemini": (BLUE_T, BLUE), "code": (GREY_T, MUTED), "people": (RED_T, RED)}[kind]
        lx = X0 + k * 2.2
        box(slide, lx, ly + 0.04, 0.22, 0.16, fill=fill, line=line, radius=0.2)
        text(slide, lx + 0.3, ly, 1.85, 0.25, [P(R(label, 9.5, F_REG, MUTED), after=0)], anchor=MSO_ANCHOR.MIDDLE)


def s_wireframes(slide):
    set_title(slide, "Wireframes/Mock diagrams of the proposed solution (optional)")
    w = 4.52
    for i, (img, cap) in enumerate([("hazard.jpg", "Hazard brief: national flood map at the home, depth gauge, flood shelters"),
                                    ("conversation.jpg", "Conversation: record, upload or type; the topics the form needs")]):
        x = X0 + i * (w + 0.18)
        picture(slide, img, x, Y0, w)
        text(slide, x, Y0 + w * 788 / 1400 + 0.08, w, 0.5, [P(R(cap, 10, F_REG, INK2), after=0)])
    text(slide, X0, 4.62, XW, 0.6, [P(R(
        "These are screens of the working prototype rather than mock-ups. All residents are fictional; flood depths and "
        "shelters are real public data.", 10.5, F_MED, INK), after=0)])


def s_architecture(slide):
    set_title(slide, "Architecture diagram of the proposed solution")
    # Left: user and the Cloud Run service.
    box(slide, X0, Y0, 1.55, 0.95, fill=RED_T, line=RED)
    text(slide, X0 + 0.08, Y0 + 0.06, 1.4, 0.85, [P(R("Care worker", 11, F_SEMI, INK), after=1),
                                                  P(R("tablet browser", 9.5, F_REG, INK2), after=0)])
    run_x, run_w = X0 + 1.85, 4.25
    box(slide, run_x, Y0, run_w, 3.95, fill=GREY_T, line=MUTED)
    text(slide, run_x + 0.12, Y0 + 0.06, run_w - 0.2, 0.3,
         [P(R("Cloud Run service (asia-northeast1), one container", 10.5, F_SEMI, INK), after=0)])
    box(slide, run_x + 0.15, Y0 + 0.42, run_w - 0.3, 0.55, fill=RGBColor(0xFF, 0xFF, 0xFF), line=GREY_L)
    text(slide, run_x + 0.25, Y0 + 0.45, run_w - 0.45, 0.5,
         [P(R("Web app: ", 10, F_SEMI, INK), R("React, Vite, Leaflet maps", 10, F_REG, INK2), after=0)],
         anchor=MSO_ANCHOR.MIDDLE)
    text(slide, run_x + 0.15, Y0 + 1.03, run_w - 0.3, 0.28,
         [P(R("FastAPI agents (Python, google-genai SDK, structured output)", 9.5, F_MED, INK2), after=0)])
    mods = [("Transcriber", "Gemini: audio to text", BLUE_T), ("Extractor", "Gemini: facts + quotes", BLUE_T),
            ("Planner", "Gemini: official form", BLUE_T), ("Critic", "rules + Gemini notes", BLUE_T),
            ("Rewriter", "Gemini: easy-read, card", BLUE_T), ("Hazard tools", "code: tiles, shelters", RGBColor(0xFF, 0xFF, 0xFF))]
    mw, mh = (run_w - 0.45) / 2, 0.66
    for i, (t, b, fill) in enumerate(mods):
        mx = run_x + 0.15 + (i % 2) * (mw + 0.15)
        my = Y0 + 1.36 + (i // 2) * (mh + 0.12)
        box(slide, mx, my, mw, mh, fill=fill, line=BLUE if fill == BLUE_T else MUTED, lw=0.75)
        text(slide, mx + 0.1, my + 0.05, mw - 0.16, mh - 0.08,
             [P(R(t, 10, F_SEMI, INK), after=0), P(R(b, 9, F_REG, INK2), after=0)], anchor=MSO_ANCHOR.MIDDLE)
    arrow(slide, X0 + 1.59, Y0 + 0.47, run_x - 0.04, Y0 + 0.69)
    # Right: managed services and public data.
    ext_x, ext_w = run_x + run_w + 0.42, X0 + XW - (run_x + run_w + 0.42)
    exts = [("Gemini API", "3.5 Flash by default; falls back to 3.1 Flash-Lite and 2.5 Flash; 3.1 Pro for review", BLUE_T, BLUE),
            ("Firestore", "Plans, edits and the audit log", GREY_T, MUTED),
            ("Secret Manager", "Gemini API key", GREY_T, MUTED),
            ("GSI public data", "Hazard Map Portal flood tiles, designated shelters, address search, base map", GREEN_T, GREEN)]
    eh, eg = 0.86, 0.17
    for i, (t, b, fill, line) in enumerate(exts):
        ey = Y0 + i * (eh + eg)
        box(slide, ext_x, ey, ext_w, eh, fill=fill, line=line, lw=0.75)
        text(slide, ext_x + 0.1, ey + 0.05, ext_w - 0.16, eh - 0.08,
             [P(R(t, 10.5, F_SEMI, INK), after=1), P(R(b, 9, F_REG, INK2), after=0)], anchor=MSO_ANCHOR.MIDDLE)
        arrow(slide, run_x + run_w + 0.04, ey + eh / 2, ext_x - 0.04, ey + eh / 2)
    text(slide, X0, Y0 + 1.15, 1.6, 2.7, [
        P(R("Facts come from code; Gemini reasons and writes; a person approves.", 10, F_MED, INK), after=6),
        P(R("Plans are drafts until a supervisor approves them.", 9.5, F_REG, INK2), after=0)])


def s_tech(slide):
    set_title(slide, "Technologies to be used in the solution")
    groups = [
        ("Google AI", ["Gemini 3.5 Flash: transcription, fact extraction, drafting, easy-read and translation",
                       "Gemini 3.1 Pro: whole-plan review (notes only)",
                       "Automatic fallback to Gemini 3.1 Flash-Lite and 2.5 Flash when a model is busy or out of quota",
                       "google-genai SDK with JSON-schema structured output"], BLUE_T, BLUE),
        ("Google Cloud", ["Cloud Run: one service for the API and web app, Tokyo region",
                          "Firestore: plans, edits and audit log",
                          "Secret Manager: the Gemini key",
                          "Cloud Build and Artifact Registry: deploy from source"], GREY_T, INK),
        ("Public data (Japan)", ["Hazard Map Portal flood-depth tiles (MLIT), maximum-scale scenario",
                                 "GSI designated emergency evacuation sites, by hazard type",
                                 "GSI address search and base-map tiles"], GREEN_T, GREEN),
        ("Application", ["Python 3.12, FastAPI, Pydantic; React, Vite, TypeScript, Leaflet",
                         "19 unit tests (pytest) for hazard reading, quote checks and safety rules",
                         "Docker image; Playwright and Gemini TTS for the demo video"], GREY_T, INK),
    ]
    w, h, gx, gy = 4.52, 1.86, 0.18, 0.2
    for i, (title, items, fill, tc) in enumerate(groups):
        x = X0 + (i % 2) * (w + gx)
        y = Y0 + (i // 2) * (h + gy)
        card(slide, x, y, w, h, title, items, fill=fill, title_color=tc, size=10, bullets=True)


def s_cost(slide):
    set_title(slide, "Estimated implementation cost (optional)")
    rows = [("Step in one plan", "Gemini calls"), ("Draft: extract facts, then fill the form", "2"),
            ("Optional whole-plan review", "1"), ("Copies: easy-read, helper card, translation", "3"),
            ("Transcribing a recorded visit (when used)", "1"), ("Most a single plan needs", "7")]
    tbl = slide.shapes.add_table(len(rows), 2, Inches(X0), Inches(Y0), Inches(5.2), Inches(0.36 * len(rows))).table
    tbl.columns[0].width, tbl.columns[1].width = Inches(3.9), Inches(1.3)
    for r, (a, b) in enumerate(rows):
        for c, val in enumerate((a, b)):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = BLUE_T if r == 0 else (GREY_T if r == len(rows) - 1 else RGBColor(0xFF, 0xFF, 0xFF))
            tf = cell.text_frame
            tf.paragraphs[0].text = ""
            run = tf.paragraphs[0].add_run()
            run.text = val
            _font(run, 10.5, F_SEMI if r in (0, len(rows) - 1) else F_REG, INK)
            tf.paragraphs[0].alignment = PP_ALIGN.RIGHT if c == 1 else PP_ALIGN.LEFT
            cell.margin_left = cell.margin_right = Inches(0.1)
    text(slide, X0, Y0 + 2.35, 5.2, 1.6, [
        P(R("A pilot of 50 residents needs at most about 350 Gemini calls.", 11, F_MED, INK), after=6),
        P(R("We will price this with the Google Cloud pricing calculator after the pilot measures real usage, rather "
            "than quote a figure now.", 10, F_REG, INK2), after=0)])
    card(slide, 5.85, Y0, X0 + XW - 5.85, 3.0, "Infrastructure", [
        "Cloud Run: one instance kept warm during evaluation (min-instances 1), scaling to zero afterwards.",
        "Firestore: one document per plan.",
        "Public GSI data is free to use with attribution, including commercially (Hazard Map Portal open-data terms).",
        "No paid third-party APIs: walking routes stay off unless a Maps key is added.",
    ], fill=GREY_T, size=10, bullets=True)


def s_snapshots(slide):
    set_title(slide, "Snapshots of the prototype")
    sw = 2.97  # two stacked shots plus captions fit above the footer
    big_w = XW - sw - 0.2
    picture(slide, "plan_top.jpg", X0, Y0, big_w)
    text(slide, X0, Y0 + big_w * 788 / 1400 + 0.06, big_w, 0.5, [P(R(
        "The draft in the official form: a source chip on every line, safety checks alongside.", 10, F_REG, INK2), after=0)])
    sx = X0 + big_w + 0.2
    sh = sw * 788 / 1400
    picture(slide, "easyread.jpg", sx, Y0, sw)
    text(slide, sx, Y0 + sh + 0.04, sw, 0.3, [P(R("Easy-read copy for the resident", 9.5, F_REG, INK2), after=0)])
    y2 = Y0 + sh + 0.38
    picture(slide, "stamped.jpg", sx, y2, sw)
    text(slide, sx, y2 + sh + 0.04, sw, 0.3, [P(R("Approved plan with the supervisor’s seal", 9.5, F_REG, INK2), after=0)])


def s_performance(slide):
    set_title(slide, "Prototype Performance report/Benchmarking")
    rows = [("Extraction check", "Result"), ("Helpers and phone numbers found", "21 of 21"),
            ("Key facts found", "58 of 60"), ("Needs (form check boxes) found", "7 of 9"),
            ("Consent read correctly", "7 of 9 *"), ("Items dropped for having no exact quote", "4 of 111")]
    tbl = slide.shapes.add_table(len(rows), 2, Inches(X0), Inches(Y0), Inches(5.3), Inches(0.36 * len(rows))).table
    tbl.columns[0].width, tbl.columns[1].width = Inches(3.9), Inches(1.4)
    for r, (a, b) in enumerate(rows):
        for c, val in enumerate((a, b)):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = BLUE_T if r == 0 else RGBColor(0xFF, 0xFF, 0xFF)
            tf = cell.text_frame
            tf.paragraphs[0].text = ""
            run = tf.paragraphs[0].add_run()
            run.text = val
            _font(run, 10.5, F_SEMI if r == 0 or c == 1 else F_REG, INK)
            tf.paragraphs[0].alignment = PP_ALIGN.RIGHT if c == 1 else PP_ALIGN.LEFT
            cell.margin_left = cell.margin_right = Inches(0.1)
    text(slide, X0, Y0 + 2.28, 5.3, 1.7, [
        P(R("Three runs over three fictional interviews we wrote, answered by Gemini 3.1 Flash-Lite after 3.5 Flash hit "
            "its free-tier quota (10 Oct 2026). A smoke test, not a benchmark.", 9.5, F_REG, INK2), after=5),
        P(R("* Both misses were safe failures: consent was set to “unclear”, which blocks approval until a person "
            "confirms it.", 9.5, F_REG, INK2), after=0)])
    sx, sw = 5.95, X0 + XW - 5.95
    stats = [("7.5–17.4 s", "per draft in browser tests (extraction and drafting)"),
             ("19 tests pass", "legend reading, quote checks, safety rules"),
             ("Hazard facts: code only", "flood depth, shelters and distances never come from the model")]
    for i, (big, small) in enumerate(stats):
        y = Y0 + i * 1.33
        box(slide, sx, y, sw, 1.18, fill=BLUE_T if i == 0 else GREY_T)
        text(slide, sx + 0.15, y + 0.1, sw - 0.3, 1.0, [P(R(big, 20, F_SEMI, BLUE if i == 0 else INK), after=1),
                                                        P(R(small, 10, F_REG, INK2), after=0)])


def s_future(slide):
    set_title(slide, "Additional Details/Future Development (if any)")
    cols = [
        ("Build next", ["Live interview copilot with the Gemini Live API, prompting for missing answers as people talk",
                        "Walking routes with the Google Routes API",
                        "Landslide and tsunami layers from the same national portal",
                        "Agent Development Kit wrapper around the agents"], GREY_T, INK),
        ("Prove it", ["Pilot with one municipality or social-welfare council: 50 residents over 8 weeks",
                      "Measure staff minutes per approved plan and the check pass rate",
                      "Test speech accuracy with older speakers and regional dialects",
                      "Expand the evaluation to 20 interviews and 30 addresses"], BLUE_T, BLUE),
        ("Production and reach", ["Japanese privacy-law (APPI) assessment, municipal network (LGWAN) access, ISMAP-listed hosting",
                                  "Export shaped for existing municipal systems",
                                  "Other register-holding systems in JAPAC; Gemini voice models also cover Hindi, "
                                  "Tamil, Vietnamese, Indonesian and Thai"], GREY_T, INK),
    ]
    w, gap = 2.95, 0.185
    for i, (title, items, fill, tc) in enumerate(cols):
        card(slide, X0 + i * (w + gap), Y0, w, 3.95, title, items, fill=fill, title_color=tc, size=11.5, bullets=True)


def s_links(slide, prototype_url):
    tb = next(s for s in slide.shapes if s.has_text_frame and "Provide links" in s.text_frame.text)
    values = {"GitHub Public Repository": GITHUB, "Demo Video Link (3 Minutes)": VIDEO,
              "Final Product Link": prototype_url}
    for p in tb.text_frame.paragraphs:
        label = p.text.strip()
        if label in values:
            r = p.runs[0]
            r.text = label + ": "
            val = values[label]
            new = copy.deepcopy(r._r)
            new.find(qn("a:t")).text = val or "Live link: added at deployment"
            r._r.addnext(new)
            run = p.runs[1]
            run.font.size = Pt(12)
            if val:
                run.hyperlink.address = val
                run.font.color.rgb = BLUE
    tb.height = Inches(1.6)
    box(slide, X0, 2.55, XW, 1.62, fill=GREY_T)
    steps = ["Open the live link and pick Hana Tanaka, the highest-risk resident.",
             "Read the hazard brief: flood depth at her home, flood shelters, the path.",
             "Start the conversation, choose “Use sample conversation”, then “Draft plan” (10 to 40 seconds).",
             "Click any source chip to see the exact words a line came from; read the checks on the right.",
             "Open “Copies and approval”: write the copies, mark as reviewed, approve, and see the seal."]
    paras = [P(R("Try it in two minutes", 12, F_SEMI, INK), after=5)]
    for i, st in enumerate(steps, 1):
        paras.append(P(R(f"{i}.  ", 11, F_SEMI, BLUE), R(st, 11, F_REG, INK2), after=3))
    text(slide, X0 + 0.15, 2.65, XW - 0.3, 2.0, paras)
    text(slide, X0, 4.32, XW, 0.4, [
        P(R("All residents in the prototype are fictional. Hazard and shelter data are real public data from GSI and "
            "the MLIT Hazard Map Portal.", 9.5, F_REG, MUTED), after=0)])


def s_responsible(slide):
    text(slide, 0.29, 0.73, 9.78, 0.61, [P(R("Responsible AI and honest limits", 16, F_MED, INK), after=0)],
         inset=0.1)
    cols = [
        ("Safety by design", ["Drafts, never decisions: a supervisor approves every plan",
                              "Facts come from code; missing data shows “no data”, never “safe”",
                              "The critic flags shelters that sit inside a flood zone",
                              "Must-fix checks block approval; every edit is logged"], BLUE_T, BLUE),
        ("Privacy", ["The prototype holds fictional residents only",
                     "Recorded audio is transcribed and not stored",
                     "Consent is recorded before a plan is shared with helpers",
                     "Plans live in Firestore in the Tokyo region; the Gemini key sits in Secret Manager"],
         GREY_T, INK),
        ("What it is not yet", ["Not yet tested with real care workers or residents, and no partner yet",
                                "Evaluation so far is a smoke test on three interviews we wrote",
                                "Japanese-language, flood-only prototype",
                                "Distances are straight-line until a Routes key is added"], RED_T, RED),
    ]
    w, gap = 2.95, 0.185
    for i, (title, items, fill, tc) in enumerate(cols):
        card(slide, X0 + i * (w + gap), Y0, w, 3.05, title, items, fill=fill, title_color=tc, size=12, bullets=True)


def build(template: Path, out: Path, prototype_url: str | None):
    prs = Presentation(str(template))
    delete_slide(prs, 1)  # the template's own "please download this template" instruction slide
    s = list(prs.slides)
    s_team(s[0])
    s_brief(s[1])
    s_opportunities(s[2])
    s_features(s[3])
    s_flow(s[4])
    s_wireframes(s[5])
    s_architecture(s[6])
    s_tech(s[7])
    s_cost(s[8])
    s_snapshots(s[9])
    s_performance(s[10])
    s_future(s[11])
    s_links(s[12], prototype_url)
    s_responsible(s[13])
    # s[14] is the template's "Thank you" slide, kept as is.
    prs.core_properties.title = "Yorisoi: AI copilot for individual evacuation plans"
    prs.core_properties.author = "The AI way"
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    return out


def to_pdf(pptx: Path) -> Path:
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(pptx.parent), str(pptx)],
                   check=True, capture_output=True, timeout=300)
    return pptx.with_suffix(".pdf")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", type=Path, default=HERE / "template.pptx")
    ap.add_argument("--out", type=Path, default=HERE / "out" / "Yorisoi-Pitch-Deck.pptx")
    ap.add_argument("--prototype-url", default=None)
    ap.add_argument("--pdf", action="store_true", help="also convert to PDF with LibreOffice")
    a = ap.parse_args()
    deck = build(a.template, a.out, a.prototype_url)
    print("wrote", deck)
    if a.pdf:
        pdf = to_pdf(deck)
        print("wrote", pdf, f"{pdf.stat().st_size / 1e6:.2f} MB")
